"""生成 8 份合成采购合同（6 DOCX + 2 PDF）与评测标签。

在 backend 目录运行：
    PYTHONUTF8=1 uv run python scripts/generate_demo_data.py

所有公司名、金额、条款均为虚构，仅用于软件测试与演示，不构成真实合同。
"""

from __future__ import annotations

import json
from pathlib import Path
from xml.sax.saxutils import escape

import yaml
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

BACKEND_DIR = Path(__file__).resolve().parents[1]
SAMPLES_DIR = BACKEND_DIR / "data" / "samples"
LABELS_DIR = BACKEND_DIR / "eval" / "labels"
FONT_PATH = Path("C:/Windows/Fonts/simhei.ttf")


def C(article: str, heading: str, *paragraphs: str) -> dict:
    return {"article": article, "heading": heading, "paragraphs": list(paragraphs)}


def L(
    article: str,
    category: str,
    severity: str,
    rule_id: str | None = None,
    keywords: tuple[str, ...] = (),
) -> dict:
    return {
        "clause_article": article,
        "category": category,
        "severity": severity,
        "rule_id": rule_id,
        "keywords": list(keywords),
    }


CONTRACTS = [
    {
        "id": "sample_01",
        "title": "数控设备采购合同",
        "filename": "sample_01_数控设备采购合同.docx",
        "format": "docx",
        "buyer": "星海智能装备有限公司",
        "supplier": "远航机电有限公司",
        "description": "数控加工中心采购，含预付款、验收、违约与管辖条款。",
        "clauses": [
            C("第一条", "合同标的", "甲方向乙方采购数控加工中心一台，型号 XH-850，具体技术参数见附件一。"),
            C("第二条", "合同价款", "合同总价为人民币 1,200,000 元（含税），包含运输、包装、安装调试费用。"),
            C(
                "第三条",
                "付款方式",
                "合同生效后 5 个工作日内，甲方向乙方支付合同总价 50% 的预付款；剩余 50% 于验收合格后 30 日内支付。",
                "乙方应在收款前开具增值税专用发票。",
            ),
            C(
                "第四条",
                "交付",
                "乙方应于合同生效后 30 日内将货物交付至甲方工厂（上海市浦东新区）。",
                "运输方式为乙方送货上门，运费由乙方承担。",
            ),
            C(
                "第五条",
                "验收",
                "甲方应在收货后 3 日内完成验收；逾期未提出书面异议的，视为验收合格。",
                "验收标准为符合国家相关标准。",
            ),
            C(
                "第六条",
                "质量保证",
                "质保期为验收合格后 12 个月；质保期内出现质量问题，乙方负责免费维修。",
            ),
            C(
                "第七条",
                "违约责任",
                "乙方逾期交付的，按逾期天数每日向甲方支付合同总价 0.01% 的违约金；违约金总额不超过合同总价的 3%。",
                "甲方逾期付款的，按同等标准向乙方支付违约金。",
            ),
            C(
                "第八条",
                "争议解决",
                "因本合同发生争议，双方应协商解决；协商不成的，提交乙方所在地人民法院诉讼解决。",
            ),
            C("第九条", "其他", "本合同自双方签字盖章之日起生效。未尽事宜，双方另行协商。"),
        ],
        "labels": [
            L("第三条", "payment", "high", "PAY-001", ("50%",)),
            L("第五条", "acceptance", "high", "ACC-002", ("3 日", "视为验收合格")),
            L("第五条", "acceptance", "medium", "ACC-003", ("国家相关标准",)),
            L("第七条", "liability", "medium", "LIA-001", ("0.01%",)),
            L("第七条", "liability", "high", "LIA-002", ("3%",)),
            L("第八条", "dispute", "high", "DIS-001", ("乙方所在地",)),
            L("合同级", "compliance", "medium", "COM-001", ()),
        ],
    },
    {
        "id": "sample_02",
        "title": "办公家具采购合同",
        "filename": "sample_02_办公家具采购合同.docx",
        "format": "docx",
        "buyer": "云杉科技有限公司",
        "supplier": "恒美办公家具有限公司",
        "description": "办公家具批量采购，含账期、质保与续约条款。",
        "clauses": [
            C("第一条", "合同标的", "甲方向乙方采购办公桌、办公椅等办公家具一批，具体名称、数量见附件清单。"),
            C("第二条", "合同价款", "合同总价为人民币 380,000 元（含税）。"),
            C(
                "第三条",
                "付款方式",
                "货物交付并验收合格后 90 日内，甲方向乙方支付全部价款。",
                "乙方应在收款前开具增值税专用发票。",
            ),
            C("第四条", "交付", "乙方应于合同生效后 20 日内送货至甲方办公楼。"),
            C("第五条", "验收", "甲方应在收货后 10 日内完成验收，验收标准为符合国家相关标准。"),
            C("第六条", "质量保证", "质保期为验收合格后 6 个月。"),
            C("第七条", "违约责任", "乙方逾期交付的，按每日合同总价的 0.02% 向甲方支付违约金。"),
            C("第八条", "合同续期", "本合同期满前，双方未提出异议的，自动续约一年。"),
            C("第九条", "争议解决", "因本合同发生争议，由甲方所在地人民法院管辖。"),
            C("第十条", "其他", "本合同自双方签字盖章之日起生效。"),
        ],
        "labels": [
            L("第三条", "payment", "medium", "PAY-003", ("90 日",)),
            L("第七条", "liability", "medium", "LIA-001", ("0.02%",)),
            L("第六条", "warranty", "high", "LIA-004", ("质量",)),
            L("第六条", "warranty", "medium", "WAR-001", ("6 个月",)),
            L("第八条", "termination", "medium", "TER-002", ("自动续约",)),
            L("第五条", "acceptance", "medium", "ACC-003", ("国家相关标准",)),
            L("合同级", "compliance", "medium", "COM-001", ()),
        ],
    },
    {
        "id": "sample_03",
        "title": "钢材年度采购合同",
        "filename": "sample_03_钢材年度采购合同.docx",
        "format": "docx",
        "buyer": "华东精密制造有限公司",
        "supplier": "鑫源钢铁贸易有限公司",
        "description": "钢材年度框架采购，含订单结算、违约责任与解除条款。",
        "clauses": [
            C(
                "第一条",
                "合同标的",
                "甲方向乙方采购钢材一批，年度预计采购量 500 吨，具体规格、数量以采购订单为准。",
            ),
            C(
                "第二条",
                "合同价款",
                "合同总价暂定为人民币 2,250,000 元（含税），按实际订单结算，单价为人民币 4,500 元/吨。",
            ),
            C("第三条", "付款方式", "甲方应在收到货物后 30 日内支付货款。"),
            C("第四条", "交付", "乙方应在收到订单后 15 日内交货至甲方仓库。"),
            C(
                "第五条",
                "验收",
                "甲方应在收货后 15 日内完成验收，验收标准见附件技术协议。",
            ),
            C("第六条", "质量保证", "质保期为验收合格后 12 个月。"),
            C(
                "第七条",
                "违约责任",
                "乙方逾期交货的，按日万分之五支付违约金；违约金累计不超过合同总价的 3%。",
                "乙方应赔偿因违约给甲方造成的全部损失，包括间接损失和可得利益损失。",
                "甲方逾期付款的，按日万分之五支付违约金。",
            ),
            C(
                "第八条",
                "合同解除",
                "乙方有权在甲方逾期付款超过 30 日时解除合同；甲方无权解除合同。",
            ),
            C(
                "第九条",
                "争议解决",
                "因本合同发生争议，提交上海仲裁委员会仲裁，也可以向甲方所在地人民法院起诉。",
            ),
            C("第十条", "其他", "本合同自双方签字盖章之日起生效。"),
        ],
        "labels": [
            L("第三条", "payment", "high", "PAY-002", ("收到货物后",)),
            L("第七条", "liability", "high", "LIA-002", ("3%",)),
            L("第七条", "liability", "high", "LIA-003", ("间接损失", "可得利益")),
            L("第六条", "warranty", "high", "LIA-004", ("质量",)),
            L("第八条", "termination", "high", "TER-001", ("甲方无权解除",)),
            L("第九条", "dispute", "high", "DIS-002", ("仲裁", "人民法院")),
            L("合同级", "compliance", "medium", "COM-001", ()),
        ],
    },
    {
        "id": "sample_04",
        "title": "仓储软件及设备采购合同",
        "filename": "sample_04_仓储软件及设备采购合同.docx",
        "format": "docx",
        "buyer": "蓝湾供应链有限公司",
        "supplier": "智仓软件科技有限公司",
        "description": "仓储管理软件及硬件采购，含部署、验收与知识产权条款。",
        "clauses": [
            C(
                "第一条",
                "合同标的",
                "甲方向乙方采购仓储管理系统软件一套及配套扫码设备 10 台，软件功能清单见附件。",
            ),
            C("第二条", "合同价款", "合同总价为人民币 680,000 元（含税）。"),
            C(
                "第三条",
                "付款方式",
                "合同生效后 5 个工作日内支付 30% 预付款；系统上线并验收合格后 90 日内支付剩余 70%。",
            ),
            C("第四条", "交付", "乙方应在合同生效后 45 日内完成软件部署和设备交付。"),
            C(
                "第五条",
                "验收",
                "甲方应在系统上线后 15 日内组织验收，验收标准为双方确认的功能清单。",
            ),
            C(
                "第六条",
                "质量保证",
                "软件免费维护期为验收合格后 12 个月；维护期内出现缺陷，乙方负责修理或更换。",
            ),
            C("第七条", "知识产权", "双方应就知识产权事宜另行签署协议。"),
            C("第八条", "保密", "乙方应对甲方的业务数据承担保密义务。"),
            C(
                "第九条",
                "违约责任",
                "乙方逾期交付的，按日万分之五支付违约金；违约金上限为合同总价的 2%。",
            ),
            C("第十条", "争议解决", "因本合同发生争议，由甲方所在地人民法院管辖。"),
            C("第十一条", "其他", "本合同自双方签字盖章之日起生效。"),
        ],
        "labels": [
            L("第三条", "payment", "high", "PAY-001", ("30%",)),
            L("第三条", "payment", "medium", "PAY-003", ("90 日",)),
            L("第七条", "ip", "medium", "IP-001", ("知识产权",)),
            L("第八条", "compliance", "medium", "COM-001", ("保密",)),
            L("第九条", "liability", "high", "LIA-002", ("2%",)),
            L("第四条", "delivery", "medium", "DEL-002", ("交付",)),
            L("第五条", "acceptance", "medium", None, ("双方确认",)),
        ],
    },
    {
        "id": "sample_05",
        "title": "食品包装材料采购合同",
        "filename": "sample_05_食品包装材料采购合同.pdf",
        "format": "pdf",
        "buyer": "绿源食品有限公司",
        "supplier": "安泰包装材料有限公司",
        "description": "食品包装材料采购，含预付款、交付与验收条款。",
        "clauses": [
            C("第一条", "合同标的", "甲方向乙方采购食品级塑料包装盒一批，规格、数量以采购订单为准。"),
            C("第二条", "合同价款", "合同总价暂定人民币 260,000 元（含税），按实际供货量结算。"),
            C(
                "第三条",
                "付款方式",
                "合同生效后 5 个工作日内支付 40% 预付款；货到后 30 日内支付剩余 60%。",
            ),
            C(
                "第四条",
                "交付",
                "乙方应在收到订单后尽快交货，具体时间另行通知。",
            ),
            C(
                "第五条",
                "验收",
                "甲方应在收货后 2 日内完成验收，逾期视为验收合格。",
            ),
            C(
                "第六条",
                "质量保证",
                "质保期为验收合格后 6 个月。",
                "质量不合格的，甲方有权要求退货或换货，费用由乙方承担。",
            ),
            C(
                "第七条",
                "违约责任",
                "乙方逾期交货的，按日万分之二支付违约金。",
                "乙方对甲方因此遭受的损失承担一切责任。",
            ),
            C(
                "第八条",
                "保密",
                "双方应对本合同内容及履行过程中获知的对方商业信息承担保密义务。",
            ),
            C("第九条", "其他", "本合同自双方签字盖章之日起生效。"),
        ],
        "labels": [
            L("第三条", "payment", "high", "PAY-001", ("40%",)),
            L("第三条", "payment", "high", "PAY-002", ("货到后",)),
            L("第四条", "delivery", "high", "DEL-001", ("尽快",)),
            L("第四条", "delivery", "medium", "DEL-002", ("交货",)),
            L("第五条", "acceptance", "high", "ACC-002", ("2 日", "视为验收合格")),
            L("第六条", "warranty", "medium", "WAR-001", ("6 个月",)),
            L("第七条", "liability", "medium", "LIA-001", ("万分之二",)),
            L("第七条", "liability", "high", "LIA-003", ("一切责任",)),
            L("合同级", "dispute", "high", "DIS-003", ()),
        ],
    },
    {
        "id": "sample_06",
        "title": "实验室仪器采购合同",
        "filename": "sample_06_实验室仪器采购合同.pdf",
        "format": "pdf",
        "buyer": "康泽生物科技有限公司",
        "supplier": "精析仪器有限公司",
        "description": "实验室仪器采购，含交付、验收与质保条款。",
        "clauses": [
            C("第一条", "合同标的", "甲方向乙方采购高效液相色谱仪一台，型号 LC-2050。"),
            C("第二条", "合同价款", "合同总价为人民币 450,000 元。"),
            C("第三条", "付款方式", "验收合格后 30 日内，甲方向乙方支付全部价款。"),
            C("第四条", "交付", "乙方应在合同生效后 60 日内交付。"),
            C(
                "第五条",
                "验收",
                "甲方应在收货后 15 日内完成验收，验收标准见附件技术参数。",
            ),
            C("第六条", "质量保证", "质保期为验收合格后 6 个月。"),
            C(
                "第七条",
                "违约责任",
                "乙方逾期交付的，按日万分之五支付违约金；违约金上限为合同总价的 10%。",
            ),
            C("第八条", "争议解决", "因本合同发生争议，由甲方所在地人民法院管辖。"),
            C("第九条", "其他", "本合同自双方签字盖章之日起生效。"),
        ],
        "labels": [
            L("第六条", "warranty", "medium", "WAR-001", ("6 个月",)),
            L("第二条", "price", "medium", "PRI-002", ("450,000",)),
            L("合同级", "compliance", "medium", "COM-001", ()),
            L("第四条", "delivery", "medium", "DEL-002", ("交付",)),
            L("第四条", "delivery", "low", None, ("运费", "风险")),
            L("第六条", "warranty", "high", "LIA-004", ("质量",)),
        ],
    },
    {
        "id": "sample_07",
        "title": "冷链保温箱采购合同",
        "filename": "sample_07_冷链保温箱采购合同.docx",
        "format": "docx",
        "buyer": "鲜达冷链有限公司",
        "supplier": "恒温箱业有限公司",
        "description": "冷链保温箱采购，含不可抗力、续约与保密条款。",
        "clauses": [
            C("第一条", "合同标的", "甲方向乙方采购冷链保温箱 200 个，规格见附件。"),
            C("第二条", "合同价款", "合同总价为人民币 160,000 元（含税）。"),
            C("第三条", "付款方式", "验收合格后 90 日内，甲方向乙方支付全部价款。"),
            C("第四条", "交付", "乙方应在合同生效后 30 日内交付至甲方仓库。"),
            C(
                "第五条",
                "验收",
                "甲方应在收货后 15 日内完成验收，验收标准为符合国家相关标准。",
            ),
            C(
                "第六条",
                "质量保证",
                "质保期为验收合格后 12 个月。",
                "质量不合格的，甲方有权要求退货、换货或修理，费用由乙方承担。",
            ),
            C(
                "第七条",
                "违约责任",
                "乙方逾期交付的，按日万分之二支付违约金；违约金累计不超过合同总价的 10%。",
            ),
            C(
                "第八条",
                "不可抗力",
                "因不可抗力导致无法履行的，乙方不承担任何责任，且无需通知甲方。",
            ),
            C("第九条", "合同续期", "本合同到期后自动续期一年。"),
            C("第十条", "保密", "乙方应对甲方的商业信息承担保密义务。"),
            C("第十一条", "争议解决", "因本合同发生争议，由乙方所在地人民法院管辖。"),
            C("第十二条", "其他", "本合同自双方签字盖章之日起生效。"),
        ],
        "labels": [
            L("第三条", "payment", "medium", "PAY-003", ("90 日",)),
            L("第七条", "liability", "medium", "LIA-001", ("万分之二",)),
            L("第八条", "liability", "medium", "FOR-001", ("不可抗力",)),
            L("第九条", "termination", "medium", "TER-002", ("自动续期",)),
            L("第十一条", "dispute", "high", "DIS-001", ("乙方所在地",)),
            L("第五条", "acceptance", "medium", "ACC-003", ("国家相关标准",)),
            L("第十条", "compliance", "medium", "COM-001", ("保密",)),
        ],
    },
    {
        "id": "sample_08",
        "title": "电子元器件采购合同",
        "filename": "sample_08_电子元器件采购合同.docx",
        "format": "docx",
        "buyer": "芯联电子有限公司",
        "supplier": "华晶元器件有限公司",
        "description": "电子元器件采购，条款相对完整，用于对照测试。",
        "clauses": [
            C(
                "第一条",
                "合同标的",
                "甲方向乙方采购电容、电阻等电子元器件一批，规格型号见附件一。",
            ),
            C("第二条", "合同价款", "合同总价为人民币 220,000 元（含 13% 增值税）。"),
            C("第三条", "付款方式", "验收合格后 30 日内，甲方向乙方支付全部价款。"),
            C("第四条", "交付", "乙方应于合同生效后 20 日内交付至甲方仓库。"),
            C(
                "第五条",
                "验收",
                "甲方应在收货后 15 日内按附件一《技术规格书》及 GB/T 2828.1 抽样标准验收。",
            ),
            C(
                "第六条",
                "质量保证",
                "质保期为验收合格后 12 个月；质保期内出现质量问题，乙方负责免费修理、更换或退货，并承担由此产生的费用。",
            ),
            C(
                "第七条",
                "违约责任",
                "乙方逾期交付的，按日万分之五支付违约金；违约金上限为合同总价的 10%。",
                "因乙方故意或重大过失造成的损失，依法承担赔偿责任。",
            ),
            C(
                "第八条",
                "合同解除",
                "乙方逾期超过 30 日或质量不合格拒不整改的，甲方有权解除合同。",
            ),
            C(
                "第九条",
                "保密",
                "双方应对因履行本合同获知的对方商业信息承担保密义务。",
            ),
            C("第十条", "争议解决", "因本合同发生争议，由甲方所在地人民法院管辖。"),
            C("第十一条", "其他", "本合同自双方授权代表签字并加盖公章之日起生效。"),
        ],
        "labels": [
            L("第三条", "payment", "low", None, ("全部价款", "质保金")),
        ],
    },
]


def set_east_asian(style, name: str = "微软雅黑", size: float = 10.5) -> None:  # type: ignore[no-untyped-def]
    style.font.name = name
    style.font.size = Pt(size)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def font_run(run, size: float, bold: bool = False, italic: bool = False) -> None:  # type: ignore[no-untyped-def]
    run.font.name = "微软雅黑"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")


def write_docx(contract: dict, path: Path) -> None:
    doc = Document()
    set_east_asian(doc.styles["Normal"])
    title = doc.add_heading(contract["title"], level=0)
    for run in title.runs:
        font_run(run, 18, bold=True)

    for line in [f"甲方：{contract['buyer']}", f"乙方：{contract['supplier']}"]:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(line)
        font_run(run, 10.5)

    for clause in contract["clauses"]:
        heading = doc.add_paragraph()
        run = heading.add_run(f"{clause['article']} {clause['heading']}")
        font_run(run, 12, bold=True)
        for text in clause["paragraphs"]:
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(text)
            font_run(run, 10.5)
            paragraph.paragraph_format.line_spacing = 1.4

    note = doc.add_paragraph()
    run = note.add_run("（本文件为合成示例，仅用于系统测试，不构成真实合同。）")
    font_run(run, 9, italic=True)
    doc.save(str(path))


def write_pdf(contract: dict, path: Path) -> None:
    if not FONT_PATH.exists():
        raise SystemExit(f"找不到中文字体：{FONT_PATH}")
    pdfmetrics.registerFont(TTFont("SimHei", str(FONT_PATH)))
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CNTitle",
        parent=styles["Title"],
        fontName="SimHei",
        fontSize=18,
        leading=26,
        alignment=1,
        spaceAfter=14,
    )
    party_style = ParagraphStyle(
        "CNParty", parent=styles["Normal"], fontName="SimHei", fontSize=10.5, leading=17
    )
    heading_style = ParagraphStyle(
        "CNHeading",
        parent=styles["Normal"],
        fontName="SimHei",
        fontSize=12,
        leading=19,
        spaceBefore=9,
        spaceAfter=3,
    )
    body_style = ParagraphStyle(
        "CNBody",
        parent=styles["Normal"],
        fontName="SimHei",
        fontSize=10.5,
        leading=17,
        wordWrap="CJK",
    )
    note_style = ParagraphStyle(
        "CNNote",
        parent=body_style,
        fontSize=8.5,
        leading=14,
        textColor="#666666",
        spaceBefore=16,
    )
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=22 * mm,
        rightMargin=22 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=contract["title"],
        author="ContractGuard synthetic data",
    )
    story = [Paragraph(escape(contract["title"]), title_style)]
    story.append(Paragraph(f"甲方：{escape(contract['buyer'])}", party_style))
    story.append(Paragraph(f"乙方：{escape(contract['supplier'])}", party_style))
    story.append(Spacer(1, 8))
    for clause in contract["clauses"]:
        story.append(
            Paragraph(
                f"{escape(clause['article'])} {escape(clause['heading'])}",
                heading_style,
            )
        )
        for text in clause["paragraphs"]:
            story.append(Paragraph(escape(text), body_style))
    story.append(Paragraph("（本文件为合成示例，仅用于系统测试，不构成真实合同。）", note_style))
    doc.build(story)


def main() -> None:
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    LABELS_DIR.mkdir(parents=True, exist_ok=True)
    samples = []
    total_labels = 0
    for contract in CONTRACTS:
        path = SAMPLES_DIR / contract["filename"]
        if contract["format"] == "pdf":
            write_pdf(contract, path)
        else:
            write_docx(contract, path)
        labels = []
        for index, label in enumerate(contract["labels"], start=1):
            labels.append({"id": f"{contract['id']}-L{index:02d}", **label})
        (LABELS_DIR / f"{contract['id']}.json").write_text(
            json.dumps(labels, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        total_labels += len(labels)
        samples.append(
            {
                "id": contract["id"],
                "title": contract["title"],
                "filename": contract["filename"],
                "format": contract["format"],
                "description": contract["description"],
            }
        )
        print(f"{contract['id']} {contract['format']:4s} labels={len(labels)} {path.name}")
    (SAMPLES_DIR / "samples.yaml").write_text(
        yaml.safe_dump(samples, allow_unicode=True, sort_keys=False, width=120),
        encoding="utf-8",
    )
    (SAMPLES_DIR / "README.md").write_text(
        "# 示例合同\n\n"
        "本目录的 8 份合同全部为**合成数据**，公司名、金额、条款均为虚构，"
        "仅用于 ContractGuard 的演示与评测，不构成真实合同或法律意见。\n\n"
        "其中 6 份为 DOCX、2 份为 PDF；运行 `scripts/generate_demo_data.py` 可重新生成。\n",
        encoding="utf-8",
    )
    print(f"done: {len(CONTRACTS)} contracts, {total_labels} labels")


if __name__ == "__main__":
    main()
