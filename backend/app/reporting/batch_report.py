"""批次合并报告：Markdown / JSON / DOCX。"""

from __future__ import annotations

import re
from copy import deepcopy
from datetime import datetime
from io import BytesIO

from docx import Document

from app.config import DISCLAIMER
from app.reporting.report import build_docx, build_markdown
from app.schemas import BatchState, JobState, JobStatus

STATUS_LABELS = {
    JobStatus.QUEUED: "排队中",
    JobStatus.RUNNING: "审查中",
    JobStatus.AWAITING_REVIEW: "待人工复核",
    JobStatus.RESUMING: "完成复核中",
    JobStatus.COMPLETED: "已完成",
    JobStatus.COMPLETED_WITH_WARNINGS: "完成（有提示）",
    JobStatus.FAILED: "失败",
}


def _demote_headings(text: str) -> str:
    lines = []
    for line in text.splitlines():
        if re.match(r"^#{1,6}\s", line):
            line = "#" + line
        lines.append(line)
    return "\n".join(lines)


def _status(value: JobStatus) -> str:
    return STATUS_LABELS.get(value, value.value)


def _overview_rows(batch: BatchState) -> list[str]:
    lines = [
        "| 合同 | 状态 | 风险 | 高 | 中 | 低 | 待复核 | 已核实高风险待复核 |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for item in batch.items:
        summary = item.summary
        lines.append(
            f"| {item.filename} | {_status(item.status)} | "
            f"{summary.total if summary else 0} | "
            f"{summary.high if summary else 0} | "
            f"{summary.medium if summary else 0} | "
            f"{summary.low if summary else 0} | "
            f"{summary.pending_review if summary else 0} | "
            f"{summary.blocking_high_pending if summary else 0} |"
        )
    return lines


def _category_rows(batch: BatchState) -> list[str]:
    if not batch.stats.by_category:
        return ["- 无"]
    lines = ["| 类别 | 总数 | 高 | 中 | 低 |", "|---|---|---|---|---|"]
    for key, bucket in sorted(batch.stats.by_category.items()):
        lines.append(
            f"| {key} | {bucket['total']} | {bucket['high']} | "
            f"{bucket['medium']} | {bucket['low']} |"
        )
    return lines


def build_batch_markdown(batch: BatchState, jobs: list[JobState]) -> str:
    stats = batch.stats
    lines = [
        "# ContractGuard 批次审查报告",
        "",
        f"> {DISCLAIMER}",
        "",
        "## 一、批次概览",
        "",
        f"- 批次 ID：{batch.batch_id}",
        f"- 生成时间：{datetime.now().astimezone().isoformat(timespec='seconds')}",
        f"- 合同数：{stats.contracts}",
        f"- 风险总数：{stats.total}（高 {stats.high} / 中 {stats.medium} / 低 {stats.low}）",
        f"- 待复核：{stats.pending_review}；已核实高风险待复核：{stats.blocking_high_pending}",
        f"- 已采纳：{stats.accepted}；已忽略：{stats.rejected}",
        f"- 状态分布：{stats.by_status}",
        "",
        "### 合同概览",
        "",
        *_overview_rows(batch),
        "",
        "### 风险类别",
        "",
        *_category_rows(batch),
        "",
    ]
    if batch.blocked_items:
        lines.extend(["## 二、未完成项（本报告不完整）", ""])
        for item in batch.blocked_items:
            lines.append(
                f"- {item.filename}（{_status(item.status)}）：{item.reason}"
            )
        lines.append("")

    lines.append("## 三、分合同报告")
    lines.append("")
    jobs_by_id = {job.job_id: job for job in jobs}
    for index, item in enumerate(batch.items, start=1):
        job = jobs_by_id.get(item.job_id)
        lines.append(f"### {index}. {item.filename}（{_status(item.status)}）")
        lines.append("")
        if job is None or job.result is None:
            lines.append("该合同没有可用的审查结果。")
            lines.append("")
            continue
        lines.append(_demote_headings(build_markdown(job.result)))
        lines.append("")
    return "\n".join(lines)


def build_batch_json(batch: BatchState, jobs: list[JobState]) -> dict:
    jobs_by_id = {job.job_id: job for job in jobs}
    contracts = []
    omitted = []
    for item in batch.items:
        job = jobs_by_id.get(item.job_id)
        if job is None or job.result is None:
            omitted.append(
                {
                    "job_id": item.job_id,
                    "filename": item.filename,
                    "status": item.status.value,
                    "reason": "没有可用的审查结果",
                }
            )
            continue
        contracts.append(
            {
                "job_id": item.job_id,
                "filename": item.filename,
                "status": item.status.value,
                "summary": item.summary.model_dump(mode="json")
                if item.summary
                else None,
                "report": job.result.model_dump(mode="json"),
            }
        )
    return {
        "batch_id": batch.batch_id,
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "stats": batch.stats.model_dump(mode="json"),
        "export_allowed": batch.export_allowed,
        "blocked_items": [
            item.model_dump(mode="json") for item in batch.blocked_items
        ],
        "contracts": contracts,
        "omitted": omitted,
        "disclaimer": DISCLAIMER,
    }


def build_batch_docx(batch: BatchState, jobs: list[JobState]) -> bytes:
    doc = Document()
    doc.add_heading("ContractGuard 批次审查报告", level=0)
    doc.add_paragraph(DISCLAIMER)
    doc.add_heading("批次概览", level=1)

    stats = batch.stats
    doc.add_paragraph(
        f"批次 ID：{batch.batch_id}；合同数：{stats.contracts}；"
        f"风险总数：{stats.total}（高 {stats.high} / 中 {stats.medium} / 低 {stats.low}）；"
        f"待复核：{stats.pending_review}；已核实高风险待复核：{stats.blocking_high_pending}"
    )
    table = doc.add_table(rows=1, cols=8)
    headers = ["合同", "状态", "风险", "高", "中", "低", "待复核", "已核实高风险待复核"]
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for item in batch.items:
        summary = item.summary
        cells = table.add_row().cells
        cells[0].text = item.filename
        cells[1].text = _status(item.status)
        cells[2].text = str(summary.total if summary else 0)
        cells[3].text = str(summary.high if summary else 0)
        cells[4].text = str(summary.medium if summary else 0)
        cells[5].text = str(summary.low if summary else 0)
        cells[6].text = str(summary.pending_review if summary else 0)
        cells[7].text = str(summary.blocking_high_pending if summary else 0)

    if batch.blocked_items:
        doc.add_heading("未完成项（本报告不完整）", level=1)
        for item in batch.blocked_items:
            doc.add_paragraph(
                f"{item.filename}（{_status(item.status)}）：{item.reason}"
            )

    jobs_by_id = {job.job_id: job for job in jobs}
    for item in batch.items:
        job = jobs_by_id.get(item.job_id)
        if job is None or job.result is None:
            continue
        doc.add_page_break()
        sub = Document(BytesIO(build_docx(job.result)))
        for element in list(sub.element.body):
            if element.tag.endswith("sectPr"):
                continue
            doc.element.body.append(deepcopy(element))

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def safe_batch_report_filename(batch_id: str, extension: str) -> str:
    safe = re.sub(r'[\\/:*?"<>|]+', "_", batch_id or "batch").strip() or "batch"
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    return f"ContractGuard_批次审查报告_{safe[:40]}_{stamp}.{extension}"
