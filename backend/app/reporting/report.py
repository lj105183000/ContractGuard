"""生成 Markdown / JSON / DOCX 审查报告。"""

from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from app.schemas import Finding, ReviewResult

SEVERITY_LABELS = {"high": "高", "medium": "中", "low": "低"}
SOURCE_LABELS = {"rule": "规则", "llm": "模型", "rule+llm": "规则+模型"}
STATUS_LABELS = {
    "pending": "待复核",
    "accepted": "已采纳",
    "rejected": "已忽略",
    "needs_human": "需人工确认",
}
TIER_LABELS = {
    "confirmed": "已核实",
    "advisory": "模型建议",
    "needs_review": "待核实",
    "unknown": "未知",
}
CATEGORY_LABELS = {
    "price": "价款",
    "payment": "付款",
    "delivery": "交付",
    "acceptance": "验收",
    "warranty": "质保",
    "liability": "违约责任",
    "termination": "解除/终止",
    "dispute": "争议解决",
    "compliance": "合规",
    "ip": "知识产权",
    "general": "其他",
}


def _rate(value: float | None) -> str:
    return f"{value:.1%}" if value is not None else "—"


def _severity(finding: Finding) -> str:
    return SEVERITY_LABELS.get(finding.severity.value, finding.severity.value)


def _source(finding: Finding) -> str:
    return SOURCE_LABELS.get(finding.source.value, finding.source.value)


def _status(finding: Finding) -> str:
    return STATUS_LABELS.get(finding.status.value, finding.status.value)


def _category(finding: Finding) -> str:
    return CATEGORY_LABELS.get(finding.category.value, finding.category.value)


def _tier(finding: Finding) -> str:
    key = finding.confidence_tier.value if finding.confidence_tier else "unknown"
    return TIER_LABELS.get(key, key)


def _finding_heading(index: int, finding: Finding) -> str:
    return f"{index}. [{_severity(finding)}] {finding.title}（{finding.clause_article or '合同级'}）"


def _route_text(report: ReviewResult) -> str:
    usage = report.usage
    if not usage:
        return ""
    route = usage.llm_route_label or usage.llm_route or ""
    models = "、".join(usage.llm_models or [])
    if route and models:
        return f"{route}（{models}）"
    return route or models


def _summary_lines(report: ReviewResult) -> list[str]:
    summary = report.summary
    tier_text = "、".join(
        f"{TIER_LABELS.get(key, key)} {value}"
        for key, value in summary.by_confidence.items()
    ) or "—"
    source_text = "、".join(
        f"{SOURCE_LABELS.get(key, key)} {value}"
        for key, value in summary.by_source.items()
    ) or "—"
    lines = [
        f"- 风险总数：{summary.total}（高 {summary.high} / 中 {summary.medium} / 低 {summary.low}）",
        f"- 置信分层：{tier_text}；来源：{source_text}",
        f"- 待复核：{summary.pending_review} 条；其中高风险 {summary.high_pending} 条"
        f"（已核实 {summary.blocking_high_pending} 条阻断导出，模型建议 {summary.advisory_high_pending} 条不阻断）",
        f"- 已采纳：{summary.accepted}；已忽略：{summary.rejected}；需人工确认：{summary.needs_human}",
        f"- 证据定位率：{_rate(summary.evidence_valid_rate)}（检查 {summary.evidence_checked} 条）；"
        f"法条 ID 可解析率：{_rate(summary.citation_valid_rate)}（检查 {summary.citation_checked} 个引用）",
        f"- 检索模式：{report.usage.embedding_mode or '—'}；"
        f"重排模式：{report.usage.rerank_mode or 'off'}",
    ]
    route_text = _route_text(report)
    route_suffix = f"；路由 {route_text}" if route_text else ""
    if report.usage.llm_calls:
        lines.append(
            f"- 模型调用：{report.usage.llm_calls} 次；"
            f"输入 {report.usage.input_tokens} tokens / 输出 {report.usage.output_tokens} tokens；"
            f"总耗时 {report.usage.duration_ms / 1000:.1f}s{route_suffix}"
        )
    else:
        lines.append(
            f"- 模型调用：0 次（{report.mode} 模式）；"
            f"总耗时 {report.usage.duration_ms / 1000:.1f}s{route_suffix}"
        )
    return lines


CN_NUMBERS = ["零", "一", "二", "三", "四", "五", "六", "七", "八", "九", "十"]


def build_markdown(report: ReviewResult) -> str:
    lines: list[str] = []
    section_no = 0

    def add_section(title: str) -> str:
        nonlocal section_no
        section_no += 1
        return f"## {CN_NUMBERS[section_no]}、{title}"

    title = report.contract.title or report.contract.filename or "采购合同"
    lines.append(f"# 合同审查报告：{title}")
    lines.append("")
    lines.append(f"> {report.disclaimer}")
    lines.append("")
    lines.append(add_section("基本信息"))
    lines.append("")
    lines.append(f"- 文件名：{report.contract.filename}")
    lines.append(f"- 文件类型：{report.contract.source_type.upper()}")
    lines.append(f"- 页数：{report.contract.page_count or '—'}；条款数：{report.contract.clause_count}")
    lines.append(f"- 审查时间：{report.generated_at}")
    lines.append(f"- 运行模式：{report.mode}")
    route_text = _route_text(report)
    if route_text:
        lines.append(f"- 模型路由：{route_text}")
    lines.append("")
    lines.append(add_section("摘要"))
    lines.append("")
    lines.extend(_summary_lines(report))
    lines.append("")

    if report.warnings:
        lines.append(add_section("运行提示"))
        lines.append("")
        for warning in report.warnings:
            lines.append(f"- {warning}")
        lines.append("")

    lines.append(add_section("风险清单"))
    lines.append("")
    if not report.findings:
        lines.append("未发现规则/模型可判定的风险。这不代表合同没有风险，请人工通读全文。")
    else:
        lines.append("| # | 等级 | 条款 | 风险 | 来源 | 状态 |")
        lines.append("|---|---|---|---|---|---|")
        for index, finding in enumerate(report.findings, start=1):
            lines.append(
                f"| {index} | {_severity(finding)} | {finding.clause_article or '合同级'} | "
                f"{finding.title} | {_source(finding)} | {_status(finding)} |"
            )
    lines.append("")

    lines.append(add_section("风险详情"))
    lines.append("")
    for index, finding in enumerate(report.findings, start=1):
        lines.append(f"### {_finding_heading(index, finding)}")
        lines.append("")
        lines.append(
            f"- 类别：{_category(finding)}；来源：{_source(finding)}；"
            f"置信：{_tier(finding)}；状态：{_status(finding)}"
        )
        if finding.clause_text:
            lines.append(f"- 条款原文：{finding.clause_text}")
        lines.append("")
        if finding.evidence_quote:
            lines.append(f"> 原文证据：{finding.evidence_quote}")
        else:
            lines.append("> 原文证据：无（合同级风险或条款缺失）")
        lines.append("")
        lines.append(f"**风险分析：** {finding.analysis or '—'}")
        lines.append("")
        lines.append("**法律依据：**")
        if finding.legal_basis_ids:
            for law_id in finding.legal_basis_ids:
                article = report.legal_basis.get(law_id)
                if article:
                    text = article.text.replace("\n", " ")
                    lines.append(f"- {article.law} {article.article}【{article.title}】：{text}")
                else:
                    lines.append(f"- {law_id}（未找到正文）")
        else:
            lines.append("- 无检索到的法条依据（模型建议）")
        lines.append("")
        if finding.rule_id or finding.rule_hint:
            rule_id = finding.rule_id or finding.rule_hint
            rule = report.playbook.get(rule_id or "")
            if rule:
                lines.append(
                    f"**企业规则：** {rule.id} {rule.title}——{rule.recommendation}"
                )
                lines.append("")
        lines.append(f"**建议修改为：** {finding.suggested_text or finding.recommendation or '—'}")
        lines.append("")
        lines.append(f"**退让方案：** {finding.fallback or '—'}")
        lines.append("")
        lines.append(f"**谈判要点：** {finding.negotiation_note or '—'}")
        lines.append("")
        if finding.review_note:
            lines.append(f"**人工意见：** {finding.review_note}")
            lines.append("")
        if finding.gate_notes:
            lines.append("**门禁提示：** " + "；".join(finding.gate_notes))
            lines.append("")

    if report.legal_basis:
        lines.append(add_section("法条附录"))
        lines.append("")
        for law_id in sorted(report.legal_basis):
            article = report.legal_basis[law_id]
            lines.append(f"### {article.law} {article.article}【{article.title}】")
            lines.append("")
            lines.append(article.text)
            lines.append("")
            lines.append(f"来源：{article.source_url}（生效日期 {article.effective_date}）")
            lines.append("")

    if report.playbook:
        lines.append(add_section("企业规则附录"))
        lines.append("")
        for rule_id in sorted(report.playbook):
            rule = report.playbook[rule_id]
            lines.append(f"### {rule.id} {rule.title}")
            lines.append("")
            lines.append(f"- 风险等级：{SEVERITY_LABELS.get(rule.severity.value, rule.severity.value)}")
            lines.append(f"- 推荐立场：{rule.recommendation}")
            lines.append(f"- 退让底线：{rule.fallback}")
            lines.append(f"- 法律依据：{', '.join(rule.legal_basis)}")
            lines.append("")

    lines.append("---")
    lines.append("")
    lines.append(report.disclaimer)
    lines.append("")
    return "\n".join(lines)


def build_json(report: ReviewResult) -> str:
    return report.model_dump_json(indent=2)


def _set_east_asian_font(style, name: str = "微软雅黑") -> None:  # type: ignore[no-untyped-def]
    style.font.name = name
    style.font.size = Pt(10.5)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def _set_run_font(run, name: str = "微软雅黑", size: float = 10.5, bold: bool = False) -> None:  # type: ignore[no-untyped-def]
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def _add_paragraph(doc: Document, text: str, bold_prefix: str | None = None):  # type: ignore[no-untyped-def]
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        _set_run_font(run, bold=True)
        rest = text[len(bold_prefix) :]
        if rest:
            run = paragraph.add_run(rest)
            _set_run_font(run)
    else:
        run = paragraph.add_run(text)
        _set_run_font(run)
    return paragraph


def _style_table(table) -> None:  # type: ignore[no-untyped-def]
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, size=9)


def build_docx(report: ReviewResult) -> bytes:
    doc = Document()
    _set_east_asian_font(doc.styles["Normal"])
    for style_name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        try:
            _set_east_asian_font(doc.styles[style_name], name="微软雅黑")
        except KeyError:
            pass

    section_no = 0

    def add_section(title: str) -> None:
        nonlocal section_no
        section_no += 1
        doc.add_heading(f"{CN_NUMBERS[section_no]}、{title}", level=1)

    title = report.contract.title or report.contract.filename or "采购合同"
    heading = doc.add_heading(f"合同审查报告：{title}", level=0)
    for run in heading.runs:
        _set_run_font(run, size=20, bold=True)
    disclaimer = doc.add_paragraph()
    run = disclaimer.add_run(report.disclaimer)
    _set_run_font(run, size=9)
    run.italic = True

    add_section("基本信息")
    for text in [
        f"文件名：{report.contract.filename}",
        f"文件类型：{report.contract.source_type.upper()}",
        f"页数：{report.contract.page_count or '—'}；条款数：{report.contract.clause_count}",
        f"审查时间：{report.generated_at}",
        f"运行模式：{report.mode}",
        *(
            [f"模型路由：{_route_text(report)}"]
            if _route_text(report)
            else []
        ),
    ]:
        _add_paragraph(doc, text)

    add_section("摘要")
    for text in _summary_lines(report):
        _add_paragraph(doc, text.lstrip("- "))

    if report.warnings:
        add_section("运行提示")
        for warning in report.warnings:
            _add_paragraph(doc, f"• {warning}")

    add_section("风险清单")
    if not report.findings:
        _add_paragraph(doc, "未发现规则/模型可判定的风险。这不代表合同没有风险，请人工通读全文。")
    else:
        table = doc.add_table(rows=1, cols=6)
        headers = ["#", "等级", "条款", "风险", "来源", "状态"]
        for index, header in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.text = header
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, bold=True)
        for index, finding in enumerate(report.findings, start=1):
            cells = table.add_row().cells
            cells[0].text = str(index)
            cells[1].text = _severity(finding)
            cells[2].text = finding.clause_article or "合同级"
            cells[3].text = finding.title
            cells[4].text = _source(finding)
            cells[5].text = _status(finding)
        _style_table(table)

    doc.add_page_break()
    add_section("风险详情")
    for index, finding in enumerate(report.findings, start=1):
        doc.add_heading(_finding_heading(index, finding), level=2)
        _add_paragraph(
            doc,
            f"类别：{_category(finding)}；来源：{_source(finding)}；状态：{_status(finding)}",
        )
        if finding.clause_text:
            _add_paragraph(doc, f"条款原文：{finding.clause_text}")
        _add_paragraph(
            doc,
            "原文证据：" + (finding.evidence_quote or "无（合同级风险或条款缺失）"),
        )
        _add_paragraph(doc, f"风险分析：{finding.analysis or '—'}")
        _add_paragraph(doc, "法律依据：")
        if finding.legal_basis_ids:
            for law_id in finding.legal_basis_ids:
                article = report.legal_basis.get(law_id)
                if article:
                    _add_paragraph(
                        doc,
                        f"• {article.law} {article.article}【{article.title}】："
                        + article.text.replace("\n", " "),
                    )
                else:
                    _add_paragraph(doc, f"• {law_id}（未找到正文）")
        else:
            _add_paragraph(doc, "• 无检索到的法条依据（模型建议）")
        if finding.rule_id or finding.rule_hint:
            rule_id = finding.rule_id or finding.rule_hint
            rule = report.playbook.get(rule_id or "")
            if rule:
                _add_paragraph(doc, f"企业规则：{rule.id} {rule.title}——{rule.recommendation}")
        _add_paragraph(
            doc,
            "建议修改为：" + (finding.suggested_text or finding.recommendation or "—"),
        )
        _add_paragraph(doc, f"退让方案：{finding.fallback or '—'}")
        _add_paragraph(doc, f"谈判要点：{finding.negotiation_note or '—'}")
        if finding.review_note:
            _add_paragraph(doc, f"人工意见：{finding.review_note}")
        if finding.gate_notes:
            _add_paragraph(doc, "门禁提示：" + "；".join(finding.gate_notes))

    if report.legal_basis:
        doc.add_page_break()
        add_section("法条附录")
        for law_id in sorted(report.legal_basis):
            article = report.legal_basis[law_id]
            doc.add_heading(f"{article.law} {article.article}【{article.title}】", level=2)
            _add_paragraph(doc, article.text)
            _add_paragraph(
                doc,
                f"来源：{article.source_url}（生效日期 {article.effective_date}）",
            )

    if report.playbook:
        add_section("企业规则附录")
        for rule_id in sorted(report.playbook):
            rule = report.playbook[rule_id]
            doc.add_heading(f"{rule.id} {rule.title}", level=2)
            _add_paragraph(doc, f"推荐立场：{rule.recommendation}")
            _add_paragraph(doc, f"退让底线：{rule.fallback}")
            _add_paragraph(doc, f"法律依据：{', '.join(rule.legal_basis)}")

    doc.add_paragraph()
    final = doc.add_paragraph()
    run = final.add_run(report.disclaimer)
    _set_run_font(run, size=9)
    run.italic = True
    final.alignment = WD_ALIGN_PARAGRAPH.LEFT

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def safe_report_filename(report: ReviewResult, extension: str) -> str:
    stem = (report.contract.title or report.contract.filename or "采购合同").strip()
    stem = "".join(char for char in stem if char not in '\\/:*?"<>|').replace(" ", "_")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    return f"ContractGuard_审查报告_{stem}_{timestamp}.{extension}"
