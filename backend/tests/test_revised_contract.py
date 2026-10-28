import re
from io import BytesIO

import pytest
from docx import Document

from app.reporting.revised_contract import (
    RevisedContractError,
    build_revised_contract,
    build_revised_docx,
    build_revised_markdown,
    safe_revised_filename,
)
from app.schemas import (
    Category,
    Clause,
    ContractMeta,
    Finding,
    FindingStatus,
    ReviewResult,
    Severity,
    Source,
)
from app.services.summary import build_summary


def _report(findings: list[Finding], mode: str = "mock") -> ReviewResult:
    parts = [
        ("cl_000", "首部", "采购合同\n甲方：A 公司；乙方：B 公司。"),
        ("cl_001", "第一条", "第一条 标的\n采购设备一套。"),
        ("cl_002", "第二条", "第二条 付款\n合同生效后支付预付款 50%。"),
        ("cl_003", "第三条", "第三条 交付\n乙方应在 30 日内交付。"),
    ]
    full_text = ""
    clauses: list[Clause] = []
    for index, (clause_id, article, text) in enumerate(parts):
        start = len(full_text)
        clauses.append(
            Clause(
                id=clause_id,
                index=index,
                article=article,
                text=text,
                start=start,
                end=start + len(text),
            )
        )
        full_text += text
        if index < len(parts) - 1:
            full_text += "\n\n"
    return ReviewResult(
        contract=ContractMeta(filename="contract.docx", title="测试合同"),
        summary=build_summary(findings),
        clauses=clauses,
        findings=findings,
        full_text=full_text,
        mode=mode,
        disclaimer="测试免责声明。",
    )


def _finding(
    finding_id: str = "f1",
    *,
    clause_id: str | None = "cl_002",
    status: FindingStatus = FindingStatus.ACCEPTED,
    source: str = "llm",
    suggested_text: str = "合同生效后 10 个工作日内支付 30% 预付款；剩余 70% 验收后 30 日内支付。",
    title: str = "预付款比例过高",
    review_note: str = "",
) -> Finding:
    return Finding(
        id=finding_id,
        clause_id=clause_id,
        clause_article="第二条" if clause_id == "cl_002" else None,
        category=Category.PAYMENT,
        severity=Severity.HIGH,
        source=Source.RULE,
        title=title,
        suggested_text=suggested_text,
        suggested_text_source=source,
        status=status,
        review_note=review_note,
    )


def test_live_llm_replaces_clause_and_preserves_heading_and_gaps():
    finding = _finding()
    report = _report([finding], mode="live")
    clause_text = report.clauses[2].text

    revised = build_revised_contract(report, job_id="job-1")

    assert revised.replaced_count == 1
    assert revised.annotated_count == 0
    assert revised.changes[0].change_type == "replace"
    assert revised.changes[0].original_text == clause_text
    assert revised.text.startswith(_report([]).full_text[:10])
    assert "\n\n第一条 标的" in revised.text
    assert "30% 预付款" in revised.text
    assert "预付款 50%" not in revised.text
    assert revised.text.count("第二条 付款") == 1
    assert revised.text.endswith("乙方应在 30 日内交付。")


def test_suggestion_with_heading_is_not_duplicated():
    finding = _finding(
        suggested_text="第二条 付款\n合同生效后 5 个工作日内支付 20% 预付款。"
    )
    report = _report([finding], mode="live")

    revised = build_revised_contract(report)

    assert revised.text.count("第二条 付款") == 1
    assert "20% 预付款" in revised.text


def test_template_suggestion_keeps_original_and_annotates():
    finding = _finding(
        source="template",
        suggested_text="将预付款比例降至 30% 以内，或要求等额预付款保函。",
    )
    report = _report([finding], mode="mock")

    revised = build_revised_contract(report)

    assert revised.replaced_count == 0
    assert revised.annotated_count == 1
    assert f"{report.clauses[2].text}\n【修订建议】" in revised.text
    assert "合同生效后支付预付款 50%" in revised.text
    assert "将预付款比例降至 30% 以内" in revised.text
    assert revised.changes[0].change_type == "annotate"
    assert revised.changes[0].original_text == report.clauses[2].text
    assert "【修订建议】" in revised.changes[0].revised_text
    assert all(
        clause.text in revised.text for clause in report.clauses
    )


def test_human_edited_suggestion_replaces_in_mock_mode():
    finding = _finding(
        source="human",
        suggested_text="第二条 付款\n合同生效后支付 10% 预付款，验收后支付 90%。",
    )
    report = _report([finding], mode="mock")

    revised = build_revised_contract(report)

    assert revised.replaced_count == 1
    assert revised.changes[0].change_type == "replace"
    assert "10% 预付款" in revised.text
    assert "预付款 50%" not in revised.text


def test_mock_mode_still_annotates_even_with_llm_source():
    finding = _finding(source="llm")
    report = _report([finding], mode="mock")

    revised = build_revised_contract(report)

    assert revised.replaced_count == 0
    assert revised.annotated_count == 1
    assert "非 live 模式" in revised.changes[0].reason
    assert "合同生效后支付预付款 50%" in revised.text


def test_multiple_findings_on_one_clause_are_all_annotated():
    first = _finding(finding_id="f1", source="llm")
    second = _finding(
        finding_id="f2",
        source="template",
        suggested_text="增加预付款保函作为收款前提。",
        title="缺少预付款担保",
    )
    report = _report([first, second], mode="live")

    revised = build_revised_contract(report)

    assert revised.replaced_count == 0
    assert revised.annotated_count == 2
    assert revised.text.count("【修订建议】") == 2
    assert "合同生效后支付预付款 50%" in revised.text
    assert "30% 预付款" in revised.text
    assert "增加预付款保函" in revised.text
    assert all(change.change_type == "annotate" for change in revised.changes)
    assert "2 条已采纳建议" in revised.changes[0].reason


def test_short_llm_suggestion_is_annotated():
    finding = _finding(suggested_text="降低比例。")
    report = _report([finding], mode="live")

    revised = build_revised_contract(report)

    assert revised.replaced_count == 0
    assert revised.annotated_count == 1
    assert "过短" in revised.changes[0].reason
    assert "合同生效后支付预付款 50%" in revised.text


def test_only_accepted_findings_are_applied():
    rejected = _finding(finding_id="f1", status=FindingStatus.REJECTED)
    pending = _finding(finding_id="f2", status=FindingStatus.PENDING)
    report = _report([rejected, pending], mode="live")

    revised = build_revised_contract(report)

    assert revised.text == report.full_text
    assert revised.changes == []
    assert revised.accepted_count == 0
    assert revised.warnings


def test_contract_level_finding_appends_supplementary_section():
    finding = _finding(
        finding_id="f1",
        clause_id=None,
        source="template",
        suggested_text="增加双向保密义务，保密期限为合同终止后三年。",
        title="缺少保密条款",
        review_note="同意补充",
    )
    report = _report([finding], mode="mock")

    revised = build_revised_contract(report)

    assert revised.appended_count == 1
    assert "补充条款（AI 建议，需人工确认）" in revised.text
    assert "1. 缺少保密条款（合同级）" in revised.text
    assert "增加双向保密义务" in revised.text
    assert "人工意见：同意补充" in revised.text
    assert revised.changes[-1].change_type == "append"
    assert revised.changes[-1].original_text == ""


def test_missing_clause_id_is_appended_with_warning():
    finding = _finding(clause_id="cl_999")
    report = _report([finding], mode="live")

    revised = build_revised_contract(report)

    assert revised.appended_count == 1
    assert "cl_999" in revised.warnings[0]
    assert "补充条款（AI 建议，需人工确认）" in revised.text


def test_invalid_span_is_appended_without_corrupting_body():
    finding = _finding(clause_id="cl_001")
    report = _report([finding], mode="live")
    mutated = report.full_text.replace("采购设备一套。", "采购设备两套。")
    report.full_text = mutated

    revised = build_revised_contract(report)

    assert revised.replaced_count == 0
    assert revised.appended_count == 1
    assert "定位不一致" in revised.warnings[0]
    assert revised.text.startswith(mutated)
    assert "采购设备一套。" not in revised.text
    assert "补充条款（AI 建议，需人工确认）" in revised.text


def test_overlapping_clauses_append_later_group():
    full_text = "AAAAABBBB"
    first = Clause(
        id="c1", index=0, article="第一条", text=full_text[0:5], start=0, end=5
    )
    second = Clause(
        id="c2", index=1, article="第二条", text=full_text[2:8], start=2, end=8
    )
    findings = [
        _finding(finding_id="f1", clause_id="c1", source="template"),
        _finding(finding_id="f2", clause_id="c2", source="template"),
    ]
    report = ReviewResult(
        contract=ContractMeta(filename="contract.txt", title="重叠测试"),
        summary=build_summary(findings),
        clauses=[first, second],
        findings=findings,
        full_text=full_text,
        mode="mock",
    )

    revised = build_revised_contract(report)

    assert revised.annotated_count == 1
    assert revised.appended_count == 1
    assert any("重叠" in warning for warning in revised.warnings)


def test_two_replacements_do_not_drift_offsets():
    first = _finding(
        finding_id="f1",
        clause_id="cl_001",
        suggested_text="第一条 标的\n采购设备两套，含安装调试。",
        title="标的约定不清",
    )
    second = _finding(
        finding_id="f2",
        clause_id="cl_002",
        suggested_text="第二条 付款\n合同生效后支付 20% 预付款，验收后支付 80%。",
    )
    report = _report([first, second], mode="live")

    revised = build_revised_contract(report)

    assert revised.replaced_count == 2
    assert revised.text.count("采购设备一套。") == 0
    assert revised.text.count("采购设备两套，含安装调试。") == 1
    assert revised.text.count("验收后支付 80%") == 1
    assert "第三条 交付\n乙方应在 30 日内交付。" in revised.text


def test_no_accepted_returns_original_with_warning():
    report = _report([], mode="mock")

    revised = build_revised_contract(report)

    assert revised.text == report.full_text
    assert revised.changes == []
    assert revised.accepted_count == 0
    assert revised.warnings == ["没有已采纳的风险，修订稿与原文一致。"]


def test_empty_full_text_raises():
    report = _report([_finding()])
    report.full_text = "   "

    with pytest.raises(RevisedContractError):
        build_revised_contract(report)


def test_markdown_and_filename():
    report = _report([_finding(source="template")], mode="mock")
    revised = build_revised_contract(report)

    markdown = build_revised_markdown(revised)

    assert "# 合同修订稿" in markdown
    assert "## 修订后合同全文" in markdown
    assert "## 修订变更清单" in markdown
    assert "【修订建议】" in markdown
    assert revised.text in markdown

    name = safe_revised_filename(report, "docx")
    assert name.endswith(".docx")
    assert "合同修订稿" in name
    assert not any(char in name for char in '/:*?"<>|')


def test_docx_contains_revised_text_and_change_list():
    report = _report([_finding(source="template")], mode="mock")
    revised = build_revised_contract(report)

    payload = build_revised_docx(revised)

    assert payload[:2] == b"PK"
    document = Document(BytesIO(payload))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "合同修订稿：测试合同" in text
    assert "修订变更清单" in text
    assert "【修订建议】" in text
    assert "原位批注" in table_text
    xml = document.element.xml
    assert not re.search(r"<w:ins[\s>/]", xml)
    assert not re.search(r"<w:del[\s>/]", xml)
